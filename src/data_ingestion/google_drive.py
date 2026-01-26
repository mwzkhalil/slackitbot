"""
Google Drive API integration for file retrieval.
"""

from googleapiclient.discovery import build
from google.oauth2 import service_account
from googleapiclient.http import MediaIoBaseDownload
import io
import logging
from typing import List, Dict, Any
from docx import Document
from pptx import Presentation
from PyPDF2 import PdfReader

class GoogleDriveAPI:
    def __init__(self, credentials_path: str):
        # Use service account credentials instead of OAuth user credentials
        creds = service_account.Credentials.from_service_account_file(
            credentials_path,
            scopes=['https://www.googleapis.com/auth/drive.readonly']
        )
        self.service = build('drive', 'v3', credentials=creds)

    async def get_files(self, folder_id: str, recursive: bool = True) -> List[Dict[str, Any]]:
        """
        Get list of files in the folder.
        If recursive=True, also fetch files from all subfolders.
        """
        all_files = []
        
        try:
            # First, verify the folder exists and we have access
            try:
                folder_info = self.service.files().get(fileId=folder_id, fields='id, name, mimeType').execute()
                logging.info(f"Accessing folder: {folder_info.get('name', 'Unknown')} (ID: {folder_id})")
            except Exception as e:
                logging.error(f"⚠️ Cannot access folder {folder_id}: {str(e)}")
                logging.error("This usually means:")
                logging.error("  1. The folder ID is incorrect")
                logging.error("  2. The service account doesn't have access to the folder")
                logging.error("  3. The folder doesn't exist")
                logging.error("Solution: Share the folder with the service account email (found in your credentials JSON file)")
                return []
            
        query = f"'{folder_id}' in parents and trashed=false"
        results = self.service.files().list(
            q=query, 
            fields="files(id, name, mimeType)",
            pageSize=1000
        ).execute()
        items = results.get('files', [])
            
            logging.info(f"Found {len(items)} items in folder (files + subfolders)")
        
        for item in items:
            if item['mimeType'] == 'application/vnd.google-apps.folder':
                # It's a folder, recursively get its contents if recursive=True
                    logging.info(f"  Found subfolder: {item['name']} (ID: {item['id']})")
                if recursive:
                    subfolder_files = await self.get_files(item['id'], recursive=True)
                    all_files.extend(subfolder_files)
                        logging.info(f"    Subfolder '{item['name']}' contains {len(subfolder_files)} files")
            else:
                # It's a file
                all_files.append(item)
                    logging.info(f"  Found file: {item['name']} (Type: {item['mimeType']})")
            
            if not items:
                logging.warning(f"⚠️ Folder '{folder_info.get('name', 'Unknown')}' appears to be empty or inaccessible")
            
        except Exception as e:
            logging.error(f"Error getting files from folder {folder_id}: {str(e)}")
            import traceback
            logging.error(traceback.format_exc())
        
        return all_files

    async def get_file_content(self, file_id: str) -> str:
        """
        Download and extract text content from file.
        Supports Google Docs, Slides, Sheets, PDFs, Word docs (.docx), PowerPoint (.pptx), and text files.
        """
        try:
            file = self.service.files().get(fileId=file_id, fields='mimeType, name').execute()
            mime_type = file['mimeType']
            file_name = file.get('name', '')

            # Handle Google Workspace files
            if mime_type == 'application/vnd.google-apps.document':
                # Export Google Doc as text
                request = self.service.files().export(fileId=file_id, mimeType='text/plain')
            elif mime_type == 'application/vnd.google-apps.presentation':
                # Export Google Slides as text
                request = self.service.files().export(fileId=file_id, mimeType='text/plain')
            elif mime_type == 'application/vnd.google-apps.spreadsheet':
                # Export Google Sheets as CSV
                request = self.service.files().export(fileId=file_id, mimeType='text/csv')
            # Handle Microsoft Office files
            elif mime_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                             'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                             'application/msword',
                             'application/vnd.ms-powerpoint']:
                # For .docx, .pptx - download and extract (basic text extraction)
                request = self.service.files().get_media(fileId=file_id)
            # Handle PDFs and other files
            else:
                # Download as-is
                request = self.service.files().get_media(fileId=file_id)

            fh = io.BytesIO()
            downloader = MediaIoBaseDownload(fh, request)
            done = False
            while done is False:
                status, done = downloader.next_chunk()

            # Try to decode as UTF-8 text
            try:
                content = fh.getvalue().decode('utf-8')
            except UnicodeDecodeError:
                # Handle binary files with proper text extraction
                fh.seek(0)  # Reset file pointer
                
                # Extract text from DOCX files
                if mime_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                'application/msword']:
                    try:
                        doc = Document(fh)
                        content = '\n'.join([para.text for para in doc.paragraphs])
                        logging.info(f"✓ Extracted {len(content)} chars from DOCX: {file_name}")
                    except Exception as e:
                        logging.error(f"✗ Error extracting DOCX {file_name}: {str(e)}")
                        content = f"[Error extracting DOCX: {str(e)}]"
                
                # Extract text from PPTX files
                elif mime_type in ['application/vnd.openxmlformats-officedocument.presentationml.presentation',
                                   'application/vnd.ms-powerpoint']:
                    try:
                        prs = Presentation(fh)
                        text_runs = []
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                if hasattr(shape, "text"):
                                    text_runs.append(shape.text)
                        content = '\n'.join(text_runs)
                        logging.info(f"✓ Extracted {len(content)} chars from PPTX: {file_name}")
                    except Exception as e:
                        logging.error(f"✗ Error extracting PPTX {file_name}: {str(e)}")
                        content = f"[Error extracting PPTX: {str(e)}]"
                
                # Extract text from PDF files
                elif mime_type == 'application/pdf':
                    try:
                        pdf = PdfReader(fh)
                        text_runs = []
                        for page in pdf.pages:
                            text_runs.append(page.extract_text())
                        content = '\n'.join(text_runs)
                        logging.info(f"✓ Extracted {len(content)} chars from PDF: {file_name}")
                    except Exception as e:
                        logging.error(f"✗ Error extracting PDF {file_name}: {str(e)}")
                        content = f"[Error extracting PDF: {str(e)}]"
                
                else:
                    content = f"[Unsupported file type: {mime_type} for file: {file_name}]"
            
            return content
        except Exception as e:
            logging.error(f"✗ Error processing file {file.get('name', 'unknown')}: {str(e)}")
            return f"[Error extracting content from file: {str(e)}]"